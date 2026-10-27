#!/usr/bin/env python
"""把已渲染的幻灯片转成 PPTX：每页为 PDF 渲染的整页高清图片（视觉与 HTML / PDF 一致，文字不可编辑），
qmd 中 ::: {.notes} 的演讲者备注写入对应页的备注栏（PowerPoint 演示者视图可见）。

用法：bash <模板>/tools/render.sh slides.qmd --pptx（推荐，一步生成 HTML / PDF / PPTX）
  或：python <模板>/tools/make_pptx.py slides.html [--pdf slides.pdf] [--width 2560] [--png] [-o slides.pptx]
  先运行 render.sh 生成 HTML 与 PDF；依赖 python-pptx、beautifulsoup4、poppler（pdftoppm）
  python 需装 python-pptx、beautifulsoup4（pip install python-pptx beautifulsoup4）
"""
import argparse
import subprocess
import sys
import tempfile
from pathlib import Path

from bs4 import BeautifulSoup
from pptx import Presentation
from pptx.util import Emu


def slide_notes(html):
    """按页序返回每页备注：reveal 的叶子 <section>（不含嵌套 section）即一页。"""
    soup = BeautifulSoup(Path(html).read_text(encoding="utf-8"), "html.parser")
    root = soup.select_one("div.reveal div.slides")
    leaves = [s for s in root.find_all("section") if not s.find("section")]
    out = []
    for s in leaves:
        a = s.find("aside", class_="notes")
        out.append(a.get_text("\n", strip=True) if a else "")
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("html")
    ap.add_argument("--pdf")
    ap.add_argument("--width", type=int, default=2560, help="每页图片宽度（像素），默认 2560")
    ap.add_argument("--png", action="store_true", help="用 PNG（无损，文件更大）；默认 JPEG q92")
    ap.add_argument("-o", "--out")
    a = ap.parse_args()
    html = Path(a.html)
    pdf = Path(a.pdf) if a.pdf else html.with_suffix(".pdf")
    out = Path(a.out) if a.out else html.with_suffix(".pptx")
    if not pdf.exists():
        sys.exit(f"找不到 {pdf}，请先运行 render.sh 导出 PDF")

    notes = slide_notes(html)
    with tempfile.TemporaryDirectory() as tmp:
        fmt = ["-png"] if a.png else ["-jpeg", "-jpegopt", "quality=92"]
        subprocess.run(["pdftoppm", *fmt, "-scale-to-x", str(a.width), "-scale-to-y", "-1", str(pdf), f"{tmp}/p"], check=True)
        pages = sorted(Path(tmp).glob("p-*"), key=lambda p: int(p.stem.split("-")[-1]))
        if len(notes) != len(pages):
            print(f"注意：HTML 有 {len(notes)} 页、PDF 有 {len(pages)} 页，备注按顺序对应，可能错位", file=sys.stderr)

        prs = Presentation()
        prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)   # 16:9（13.333 × 7.5 英寸）
        blank = prs.slide_layouts[6]
        for i, img in enumerate(pages):
            s = prs.slides.add_slide(blank)
            s.shapes.add_picture(str(img), 0, 0, width=prs.slide_width, height=prs.slide_height)
            if i < len(notes) and notes[i]:
                s.notes_slide.notes_text_frame.text = notes[i]
        prs.save(out)
    n_notes = sum(1 for x in notes[: len(pages)] if x)
    print(f"已生成 {out}：{len(pages)} 页，其中 {n_notes} 页有备注，{out.stat().st_size / 1e6:.1f} MB")


if __name__ == "__main__":
    main()
